"""
Builds store/app/seed/catalogue.json — the products loaded when the website is installed.

Usage:  python3 tools/build_seed.py <Setwel-Product-Catalogue.xlsx>

Only SELLING prices go into the file (the repository is public, so supplier costs stay out).
Pricing used (same as the default Pricing rules in the store):
  - HP, Brother and Canon ink & toner: supplier price, no markup
  - Canon printers: +55%
  - everything else: +45%
  - always rounded UP to the next rand
"""
import json
import math
import re
import sys

import openpyxl

SRC = sys.argv[1]
OUT = 'store/app/seed/catalogue.json'

CATEGORY = {
    'Printers': 'Printers & Scanners', 'Scanners': 'Printers & Scanners',
    'Ink & Toner': 'Ink & Toner',
    'Labelling Machines': 'Labelling Machines & Tape',
    'USB Flash Drives': 'USB, SSD & HDD', 'SSD Drives': 'USB, SSD & HDD', 'Hard Drives': 'USB, SSD & HDD', 'Memory Cards': 'USB, SSD & HDD',
    'Batteries': 'Accessories', 'Power Banks': 'Accessories', 'Accessories': 'Accessories', 'Laptop Bags': 'Accessories',
    'Cleaning Products': 'Cleaning & Hygiene',
    'Laptops': 'Laptops, Monitors & Projectors', 'Monitors': 'Laptops, Monitors & Projectors', 'Projectors': 'Laptops, Monitors & Projectors',
}


def markup(brand, category):
    if category == 'Ink & Toner' and brand in ('HP', 'Brother', 'Canon'):
        return 0.0
    if category == 'Printers & Scanners' and brand == 'Canon':
        return 0.55
    return 0.45


wb = openpyxl.load_workbook(SRC, data_only=True)
new_skus = {str(r[0]) for r in wb['New in Market'].iter_rows(min_row=5, values_only=True) if r[0]}
ws = wb['Website import']
head = [c.value for c in ws[1]]
col = {h: i for i, h in enumerate(head)}
items, unknown = [], set()
for row in ws.iter_rows(min_row=2, values_only=True):
    if not row[0]:
        continue
    g = lambda h: row[col[h]] if row[col[h]] is not None else ''
    cat = CATEGORY.get(g('Category'))
    if not cat:
        unknown.add(g('Category'))
        continue
    cost = g('Supplier cost (excl VAT)')
    price = math.ceil(round(float(cost) * (1 + markup(g('Brand'), cat)), 6)) if cost not in ('', None) else None
    specs = '\n'.join(s.strip() for s in str(g('Specifications')).split(' | ') if s.strip())
    compatible = '\n'.join(s.strip() for s in re.split(r';|\n', str(g('Compatible printers'))) if s.strip())
    items.append({
        'sku': str(row[0]), 'name': g('Name'), 'brand': g('Brand'), 'category': cat, 'price': price,
        'short': g('Short description'), 'description': g('Description'), 'specs': specs, 'compatible': compatible,
        'mpn': str(g('MPN')), 'gtin': str(g('GTIN')), 'warranty': g('Warranty'), 'is_new': str(row[0]) in new_skus,
    })
if unknown:
    raise SystemExit(f'Unmapped categories: {unknown}')
with open(OUT, 'w', encoding='utf-8') as f:
    json.dump(items, f, ensure_ascii=False, indent=0)
print(f'Wrote {OUT}: {len(items)} products, {sum(i["is_new"] for i in items)} new in market, '
      f'{sum(1 for i in items if i["price"] is None)} price on request')
from collections import Counter
for k, v in sorted(Counter(i['category'] for i in items).items()):
    print(f'  {k:32} {v}')
