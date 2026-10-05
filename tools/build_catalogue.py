"""
Builds the Setwel Africa product catalogue workbook from a supplier price list (.xlsx).

Usage:  python3 tools/build_catalogue.py <supplier-price-list.xlsx> <output.xlsx>

What it does
- Picks the product groups Setwel asked for (cartridges per brand, cleaning, Verbatim,
  Logitech, Duracell, labelling machines, SanDisk, and a "New in the market" list).
- Skips items the supplier marks EOL (end of life).
- Cleans up the product names, pulls specs (colour, page yield, capacity, compatible
  printers) out of the supplier description, and writes a short description.
- Selling price = supplier cost (excl. VAT) x (1 + markup), rounded up to the next rand.
  The markup is one cell on the Summary sheet, so every price follows it.
- Adds a "Website import" sheet in the exact format of the store's importer.
"""
import math
import re
import sys
from collections import OrderedDict

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

SRC = sys.argv[1]
OUT = sys.argv[2]
MARKUP = 0.55

# ---------------------------------------------------------------- read supplier list
wb_in = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
ws_in = wb_in.worksheets[0]
def to_price(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None  # e.g. "P.O.R" (price on request)


ROWS = []
for r in ws_in.iter_rows(min_row=3, values_only=True):
    if not r or not r[4] or r[7] is None:
        continue
    flag = (r[0] or '').strip().upper()
    ROWS.append({
        'flag': flag, 'group': r[2] or '', 'cat': r[3] or '', 'code': str(r[4]).strip(),
        'mpn': str(r[5] or '').strip(), 'desc': re.sub(r'\s+', ' ', str(r[6] or '')).strip().strip('"'),
        'cost': to_price(r[7]), 'barcode': str(r[9] or '').strip(),
    })
BY_CODE = {r['code']: r for r in ROWS}
LIVE = [r for r in ROWS if r['flag'] != 'EOL']


def pick(codes):
    out = []
    for c in codes:
        r = BY_CODE.get(c)
        if r is None:
            raise SystemExit(f'Code {c} not found in price list')
        if r['flag'] == 'EOL':
            raise SystemExit(f'Code {c} is EOL')
        out.append(r)
    return out


# ---------------------------------------------------------------- text clean-up
FIXES = [
    (r'\bHp\b', 'HP'), (r'\bHpe\b', 'HPE'), (r'\bLaserjet\b', 'LaserJet'), (r'\bOfficejet\b', 'OfficeJet'),
    (r'\bDeskjet\b', 'DeskJet'), (r'\bPagewide\b', 'PageWide'), (r'\bMfp\b', 'MFP'), (r'\bSfp\b', 'SFP'),
    (r'\bLj\b', 'LaserJet'), (r'\bEnt\b', 'Enterprise'), (r'\bLbp', 'LBP'), (r'\bMf(?=\d)', 'MF'),
    (r'\bCrg[- ]?', 'CRG-'), (r'\bGi-', 'GI-'), (r'\bPgi-?', 'PGI-'), (r'\bPg-', 'PG-'), (r'\bCl-', 'CL-'),
    (r'\bEcosys\b', 'ECOSYS'), (r'\bTaskalfa\b', 'TASKalfa'), (r'\bUsb\b', 'USB'), (r'\bUsb-c\b', 'USB-C'), (r'\bUsb-a\b', 'USB-A'),
    (r'\bSsd\b', 'SSD'), (r'\bHdd\b', 'HDD'), (r'\bTb\b', 'TB'), (r'\bGb\b', 'GB'), (r'\bMah\b', 'mAh'),
    (r'\bPixma\b', 'PIXMA'), (r'\bMaxify\b', 'MAXIFY'), (r'\bTn-', 'TN-'), (r'\bDr-', 'DR-'), (r'\bLc-?(?=\d)', 'LC'),
    (r'\bTze-', 'TZe-'), (r'\bPt-', 'PT-'), (r'\bP-touch\b', 'P-touch'), (r'\bEvomore\b', 'EvoMore'),
    (r'\bXl\b', 'XL'), (r'\bMfc[- ]?(?=[A-Za-z]?\d|[A-Za-z]\d)', 'MFC-'), (r'\bDcp[- ]?(?=[A-Za-z]?\d)', 'DCP-'),
    (r'\bHl[- ]?(?=[A-Za-z]?\d)', 'HL-'), (r'\bHll(?=\d)', 'HL-L'), (r'\bMfcl(?=\d)', 'MFC-L'), (r'\bDcpl(?=\d)', 'DCP-L'), (r'\bTri:colour\b', 'Tri-colour'), (r'\bAa\b', 'AA'), (r'\bAaa\b', 'AAA'),
    (r'\bSandisk\b', 'SanDisk'), (r'\bMicrosd\b', 'microSD'), (r'\bMicrosdxc\b', 'microSDXC'), (r'\bSdxc\b', 'SDXC'),
    (r'\bSdhc\b', 'SDHC'), (r'\bWi-fi\b', 'Wi-Fi'), (r'\bWifi\b', 'Wi-Fi'), (r'\bHd\b', 'HD'), (r'\bFhd\b', 'FHD'),
]


def fix_case(s):
    s = s.replace(':', '/').replace(' ;', ';').replace('´', "'").replace('‘', "'").replace('’', "'")
    s = re.sub(r'\s*/\s*', '/', s)
    for pat, rep in FIXES:
        s = re.sub(pat, rep, s, flags=re.I)
    # model codes such as 145a, w1450a, m2040dn -> upper case (letters+digits mixed, short)
    s = re.sub(r'\b(?=[A-Za-z]*\d)(?=\d*[A-Za-z])[A-Za-z0-9-]{2,12}\b',
               lambda m: m.group(0).upper() if not re.fullmatch(r'\d+(ml|l|kg|g|mm|cm|w|v|mah|gb|tb|mb|pk|pack|s|pgs|k)', m.group(0), re.I) else m.group(0), s)
    s = re.sub(r'(\d)\s?(gb|tb|mb)\b', lambda m: m.group(1) + m.group(2).upper(), s, flags=re.I)
    s = re.sub(r'(\d)\s?l\b', r'\1L', s)
    s = re.sub(r'(\d)ML\b', r'\1ml', s)
    # printer model suffixes read better in lower case: M454DN -> M454dn, MZ3200I -> MZ3200i
    s = re.sub(r'\b([A-Z]{0,4}-?\d{2,5})(DN|DW|CDW|CDN|IDN|CI|CIDN|FDN|FDW|NW|DNW|CIDW|FDN|I|W|CW|ADN|ADW|SDN|E)\b',
               lambda m: m.group(1) + m.group(2).lower(), s)
    s = re.sub(r'\b(\d{3})E (EvoMore)', r'\1e \2', s)
    s = re.sub(r'^Logitec\s', '', s)
    s = re.sub(r'(\d)(ML|KG|G)\b', lambda m: m.group(1) + m.group(2).lower(), s)
    s = re.sub(r'\bAll/in/one\b', 'All-in-One', s, flags=re.I)
    s = re.sub(r'\bHP(\w)', r'HP \1', s) if s.startswith('HPHP') else s
    return re.sub(r'\s{2,}', ' ', s).strip(' ,;-/')


COLOURS = ['Tri-colour', 'Photo Black', 'Matte Black', 'Matt Black', 'Black', 'Cyan', 'Magenta', 'Yellow', 'Grey', 'Red',
           'Blue', 'Green', 'Pink', 'Silver', 'White', 'Orange', 'Purple', 'Navy']


def colour_of(text):
    t = text.lower().replace('tri:colour', 'tri-colour').replace('mgenta', 'magenta').replace('yellowtoner', 'yellow toner')
    for c in COLOURS:
        if re.search(r'\b' + c.lower() + r'\b', t):
            return c
    return ''


def page_yield(text):
    t = text.replace(' ', ' ')
    num = r'(\d{1,3}(?: \d{3})+|\d{3,6})'
    m = re.search(num + r'\s*(?:page|pages|pgs)\b', t, re.I) or \
        re.search(r'(?:yield|yeild)(?:\s*of)?\s*[:-]?\s*' + num, t, re.I) or \
        re.search(r'approx\.?\s*' + num, t, re.I) or re.search(r'(\d{1,3})k\s*yield', t, re.I)
    if not m:
        return ''
    n = int(m.group(1).replace(' ', '')) * (1000 if re.search(r'k\s*yield', m.group(0), re.I) else 1)
    return f'{n:,}'.replace(',', ' ') if 50 <= n <= 500000 else ''


def compatible_of(text):
    m = re.search(r'\bfor\b\s+(.*)', text, re.I)
    if not m:
        return ''
    c = re.split(r'(?:page\s*yield|\(|\byield\b|\byeild\b|page\s*y|\bapprox)', m.group(1), flags=re.I)[0]
    c = fix_case(c)
    c = re.sub(r'(\s+\d{1,3}(?: \d{3})+|\s+\d{3,6}\s*(pgs|pages|k))$', '', c, flags=re.I)
    c = re.sub(r'(?<=[A-Za-z])\d{3,6}\s*pgs$', '', c, flags=re.I)
    c = re.sub(r'\s+(Only|Series)$', lambda m: '' if m.group(1) == 'Only' else m.group(0), c).strip(' ,;-/')
    return c if 3 <= len(c) <= 200 else ''


def capacity_of(text):
    m = re.search(r'(\d+(?:\.\d+)?)\s*(TB|GB|mAh)\b', text, re.I)
    return f'{m.group(1)}{m.group(2).upper() if m.group(2).lower() != "mah" else "mAh"}' if m else ''


def short_name(r, brand):
    """Product name: supplier description up to ' For ' / yield text, cleaned."""
    d = r['desc']
    d = re.split(r'\bfor\b', d, flags=re.I)[0] if len(d) > 70 or ' for ' in d.lower() else d
    d = re.split(r'(?:page\s*y(?:ie|ei)ld|\(\s*\d|(?<!high )(?<!xl )\byield\b|\s-\s*yield|approx)', d, flags=re.I)[0]
    d = fix_case(d)
    if len(d) > 90:  # long marketing text: keep the first clause
        d = re.split(r'(?<=[a-z0-9)])\s(?:is|are|offers|features|designed|delivers)\b', d)[0]
        d = re.sub(r'^The\s+', '', d)
    d = re.sub(r'(\s+(\d{1,3}(?: \d{3})+|\d{4,6})(\s*(pgs|pages?))?)+$', '', d, flags=re.I)
    d = re.sub(r'/[A-Za-z]{1,3}$', '', d)
    if len(d) > 60:
        cut = re.split(r'\s(?:Strong|Used|Concentrated|20%|All Purpose|Water|Thick(?:ened)?|Ready|An Acid|Liquid Sulphuric|For|Kills|Contains|Fast|Round Compressed|Removes|Deodorizes|High Foaming|Qwerty|Handheld 2)\b', d)[0]
        d = cut if len(cut) >= 12 else d[:60]
    if len(d) > 45 and ';' in d:
        d = d.split(';')[0]
    if len(d) > 45:
        keep, small = [], {'and', 'in', 'with', 'for', 'of', 'x', '&', '-'}
        for w in d.split(' '):
            if keep and w[:1].islower() and w not in small and not re.search(r'\d', w):
                break
            keep.append(w)
        while keep and keep[-1].lower() in small | {'is', 'the'}:
            keep.pop()
        d = ' '.join(keep)
    d = re.sub(r'\s(Used|For|To|With|And|Strong|Concentrated|20%.*)$', '', d.strip(' :-,'), flags=re.I)
    d = d.strip(' :-,;')
    if brand and not d.lower().startswith(brand.lower()):
        d = brand + ' ' + d
    return d.strip()[:110]


def brand_of(r):
    first = r['desc'].split(' ')[0].strip('"').lower()
    known = {'hp': 'HP', 'hpe': 'HPE', 'canon': 'Canon', 'kyocera': 'Kyocera', 'ricoh': 'Ricoh', 'brother': 'Brother',
             'pantum': 'Pantum', 'verbatim': 'Verbatim', 'logitech': 'Logitech', 'duracell': 'Duracell', 'sandisk': 'SanDisk',
             'computacare': 'Computacare', 'multipro': 'Multipro', 'addis': 'Addis', 'goldenmarc': 'Goldenmarc', 'genius': 'Genius',
             'volkano': 'Volkano', 'dymo': 'Dymo', 'kodak': 'Kodak', 'asus': 'Asus', 'viewsonic': 'ViewSonic', 'novaro': 'Novaro',
             'epson': 'Epson', 'casio': 'Casio', 'the': None}
    if first in known and known[first]:
        return known[first]
    m = re.search(r'\b(Brother|Dymo|Logitech|Kodak|Epson|Casio|Duracell|SanDisk|Verbatim|Pantum)\b', r['desc'], re.I)
    if m:
        return known.get(m.group(1).lower(), m.group(1))
    return {'BROTHER CONSUMABLES': 'Brother', 'LOGITECH': 'Logitech', 'DURACELL': 'Duracell', 'SANDISK': 'SanDisk',
            'PANTUM CONSUMABLES': 'Pantum'}.get(r['cat'], r['group'] if r['group'] not in ('OEM', 'Accessories', 'Kolok Brands') else '')


def cartridge_type(text):
    t = text.lower()
    for k, v in [('drum', 'Drum unit'), ('imaging unit', 'Imaging unit'), ('reload kit', 'Toner reload kit'), ('recharge kit', 'Toner reload kit'),
                 ('printhead', 'Printhead'), ('maintenance', 'Maintenance kit'), ('waste', 'Waste toner box'), ('collection unit', 'Waste toner box'),
                 ('belt', 'Transfer belt'), ('fuser', 'Fuser'), ('gel', 'Gel cartridge'), ('bottle', 'Ink bottle'), ('ink tank', 'Ink tank'),
                 ('ink', 'Ink cartridge'), ('ribbon', 'Ribbon'), ('tape', 'Label tape'), ('toner', 'Toner cartridge'), ('print cart', 'Toner cartridge')]:
        if k in t:
            return v
    return 'Cartridge'


def describe(item):
    """One short description built only from facts in the supplier data."""
    b, typ, col, yld, comp, cap = item['brand'], item['type'], item['colour'], item['yield'], item['compatible'], item['capacity']
    if item['kind'] == 'cartridge':
        parts = [f'Genuine {b} {col.lower() + " " if col else ""}{typ.lower()}'.replace('  ', ' ')]
        if comp:
            parts.append(f'for {comp}')
        s = ' '.join(parts) + '.'
        if yld:
            s += f' Approx. {yld} pages.'
        return s
    if cap:
        return f'{item["name"]} — {cap} capacity.'
    return item['name'] + '.'


def build_item(r, kind, category, typ=None, brand=None):
    brand = brand or brand_of(r)
    name = short_name(r, brand)
    item = OrderedDict(
        sku=r['code'], mpn=r['mpn'], barcode=r['barcode'], brand=brand, category=category, kind=kind,
        type=typ or (cartridge_type(r['desc']) if kind == 'cartridge' else ''),
        name=name, colour=colour_of(r['desc']), yield_='', compatible='', capacity=capacity_of(r['desc']),
        cost=round(r['cost'], 2) if r['cost'] is not None else None, source=r['desc'], new=r['flag'] == 'NEW',
    )
    if kind == 'cartridge':
        model = ('TK-' + r['code'][2:]) if brand == 'Kyocera' and r['code'].startswith('TK') else (r['mpn'] if brand in ('Brother', 'Pantum', 'Ricoh') else '')
        if model:
            stem = re.match(r'[A-Z]+-?\d+', model.upper())
            flat = re.sub(r'[\s-]', '', name.upper())
            if stem and stem.group(0).replace('-', '') not in flat:
                name = re.sub(r'^(' + re.escape(brand) + r')\s+', r'\1 ' + model + ' ', name)
                item['name'] = name
        item['yield_'] = page_yield(r['desc'])
        item['compatible'] = compatible_of(r['desc'])
        kn = re.search(r'\s(\d+(?:\.\d)?)k$', item['name'], re.I)
        if kn and not item['yield_']:
            item['yield_'] = f"{int(float(kn.group(1)) * 1000):,}".replace(',', ' ')
        item['name'] = re.sub(r'\s\d+(\.\d)?k$', '', item['name'], flags=re.I)
        m = re.match(r'(.*\b(?:Cartridge|Drum Unit|Drum Set|Drum|Toner|Ink Bottle|Ink|Ribbon))\s+(.+)$', item['name'])
        if m and not item['compatible'] and re.search(r'\d', m.group(2)) and not re.fullmatch(r'[\d.,]+k', m.group(2), re.I) and not re.match(r'(' + '|'.join(COLOURS) + r'|High|Standard|Original|Kit|Cartridge|Reload)\b', m.group(2), re.I):
            item['name'], item['compatible'] = m.group(1), m.group(2).strip(' /;,')
        kc = re.fullmatch(r'(\d+(?:\.\d)?)k', item['compatible'], re.I)
        if kc:  # e.g. Ricoh "18K" = 18 000 pages, not a printer
            if not item['yield_']:
                item['yield_'] = f"{int(float(kc.group(1)) * 1000):,}".replace(',', ' ')
            item['compatible'] = ''
        y = item['yield_'].replace(' ', '')
        if y:  # drop the page-yield number the supplier left at the end of the printer list / name
            item['compatible'] = re.sub(r'\s+' + r'\s?'.join(y) + r'(\s*(pgs|pages?|page\s*y(?:ie|ei)ld))?$', '', item['compatible'], flags=re.I).strip(' ,;/')
            item['name'] = re.sub(r'\s+' + r'\s?'.join(y) + r'(\s*(pgs|pages?|page yei?ld))?$', '', item['name'], flags=re.I).strip()
    item['yield'] = item['yield_']
    item['short'] = describe(item)
    if kind == 'general':
        full = fix_case(re.sub(r'^(The|he)\s+', '', r['desc']))
        item['description'] = full if len(full) > len(item['name']) + 15 else ''
    flags = []
    if r['cost'] is None:
        flags.append('Price on request (no supplier price)')
    if not r['desc']:
        flags.append('No supplier description — add name & description')
    item['check'] = '; '.join(flags)
    specs = []
    if item['type']:
        specs.append(('Type', item['type']))
    if item['colour']:
        specs.append(('Colour', item['colour']))
    if item['yield']:
        specs.append(('Page yield (approx.)', item['yield'] + ' pages'))
    if item['capacity']:
        specs.append(('Capacity', item['capacity']))
    if item['mpn']:
        specs.append(('Manufacturer code', item['mpn']))
    item['specs'] = specs
    if item['sku'] in NAME_OVERRIDES:
        item['name'] = NAME_OVERRIDES[item['sku']]
    return item


NAME_OVERRIDES = {
    'MDR3355': 'Brother DR3355 Drum Unit',
    'PTL410H': 'Pantum TL410H Black Toner Cartridge',
    'M32251': "Verbatim Charge 'N' Go 2C Magnetic Wireless Power Bank 5000mAh — Grey",
    'M32253': "Verbatim Charge 'N' Go 2C Magnetic Wireless Power Bank 5000mAh — Gold",
    'M32276': "Verbatim Charge 'N' Go 2C Magnetic Wireless Power Bank 5000mAh — Silver",
    'MPTM95AD': 'Brother P-touch PT-M95 Handheld Label Printer',
}


# ---------------------------------------------------------------- selections
SHEETS = OrderedDict()

hp_toner_codes = ['HW1450A', 'HW1450X', 'HW1510A', 'HW1510X', 'HW1530A', 'HW1530X', 'HW1540A', 'HW1540X',
                  'HW2220A', 'HW2220X', 'HW2221A', 'HW2221X', 'HW2222A', 'HW2222X', 'HW2223A', 'HW2223X',
                  'HW2300A', 'HW2300X', 'HW2301A', 'HW2301X', 'HW2302A', 'HW2302X', 'HW2303A', 'HW2303X',
                  'HW2250A', 'HW2251A', 'HW2252A', 'HW2253A',
                  'HW1360A', 'HW1360X', 'HW1500A', 'HW2130A', 'HW2131A', 'HW2132A', 'HW2133A']
hp_ink_codes = ['H4K0V9PE', 'H4K0V6PE', 'H4K0V7PE', 'H4K0V8PE', 'H4K0W3PE', 'H4K0W0PE', 'H4K0W1PE', 'H4K0W2PE',
                'H4S6X8PE', 'H4S6X5PE', 'H4S6X6PE', 'H4S6X7PE', 'H4S6Y2PE', 'H3YM62AE', 'H3YM63AE']
canon_toner_codes = ['CCRG075BK', 'CCRG075C', 'CCRG075M', 'CCRG075Y', 'CCRG075HBK', 'CCRG075HC', 'CCRG075HM', 'CCRG075HY',
                     'CCRG071', 'CCRG071H', 'CCRG070BK', 'CCRG070HBK', 'CCRG069BK', 'CCRG069HBK', 'CCRG069C']
canon_ink_codes = ['CGI46BK', 'CGI46C', 'CGI46M', 'CGI46Y', 'CGI45BK', 'CGI45C', 'CGI45M', 'CGI45Y',
                   'CGI43BK', 'CGI43C', 'CGI43M', 'CGI43Y', 'CGI43R', 'CGI43GY', 'CGI41SBK']
kyocera_codes = ['TK3400', 'TK7135', 'TK8375K', 'TK8375C', 'TK8375M', 'TK8375Y', 'TK5270K', 'TK5270M', 'TK5270Y', 'TK3190']

INK = 'Ink & Toner'
SHEETS['Brother'] = ('Brother — all cartridges & consumables', [build_item(r, 'cartridge', INK, brand='Brother') for r in LIVE if r['cat'] == 'BROTHER CONSUMABLES'])
SHEETS['HP Toner'] = ('HP — 35 newest toner cartridges', [build_item(r, 'cartridge', INK, brand='HP') for r in pick(hp_toner_codes)])
SHEETS['HP Ink'] = ('HP — 15 newest ink cartridges', [build_item(r, 'cartridge', INK, brand='HP') for r in pick(hp_ink_codes)])
SHEETS['Canon Toner'] = ('Canon — 15 newest toner cartridges', [build_item(r, 'cartridge', INK, brand='Canon') for r in pick(canon_toner_codes)])
SHEETS['Canon Ink'] = ('Canon — 15 newest ink bottles & cartridges', [build_item(r, 'cartridge', INK, brand='Canon') for r in pick(canon_ink_codes)])
SHEETS['Kyocera'] = ('Kyocera — 10 newest toners', [build_item(r, 'cartridge', INK, brand='Kyocera') for r in pick(kyocera_codes)])
SHEETS['Ricoh'] = ('Ricoh — all cartridges & consumables', [build_item(r, 'cartridge', INK, brand='Ricoh') for r in LIVE if 'ricoh' in r['desc'].lower() and r['cat'] == 'SUB D CONSUMABLES & ACCESSORIES'])
SHEETS['Pantum'] = ('Pantum — all cartridges & consumables', [build_item(r, 'cartridge', INK, brand='Pantum') for r in LIVE if r['cat'] == 'PANTUM CONSUMABLES'])

cleaning = [r for r in LIVE if r['cat'] == 'CLEANING PRODUCTS'] + pick(['GCK303', 'VK5091CL'])
SHEETS['Cleaning'] = ('Cleaning — everything for cleaning', [build_item(r, 'general', 'Cleaning Products', 'Cleaning') for r in cleaning])


def verbatim_cat(r):
    d = r['desc'].lower()
    if r['cat'] == 'VERBATIM BATTERIES':
        return 'Batteries', 'Battery'
    if r['cat'] == 'VERBATIM USBS':
        return 'USB Flash Drives', 'USB flash drive'
    if r['cat'] == 'VERBATIM HARD DRIVES':
        return 'Hard Drives', 'Hard drive'
    if 'power bank' in d and 'monitor' not in d:
        return 'Power Banks', 'Power bank'
    return None


verb = []
for r in LIVE:
    if r['group'] == 'Verbatim':
        vc = verbatim_cat(r)
        if vc:
            verb.append(build_item(r, 'general', vc[0], vc[1], 'Verbatim'))
order = {'Battery': 0, 'USB flash drive': 1, 'Hard drive': 2, 'Power bank': 3}
verb.sort(key=lambda i: (order[i['type']], i['cost'] or 0))
SHEETS['Verbatim'] = ('Verbatim — batteries, USB drives, hard drives, power banks', verb)


def logitech_type(d):
    d = d.lower()
    has_keys = bool(re.search(r'keyboard|\bkeys\b|keycap|f-keys|hot keys', d))
    has_mouse = bool(re.search(r'mouse|optical tracking|dpi|clicks life', d))
    if 'combo' in d or (has_keys and has_mouse):
        return 'Keyboard & mouse combo'
    if has_keys:
        return 'Keyboard'
    if 'presenter' in d or 'laser pointer' in d:
        return 'Presenter'
    if has_mouse:
        return 'Mouse'
    if 'webcam' in d or re.search(r'\b(720|1080)p', d):
        return 'Webcam'
    if 'headset' in d or 'headband' in d:
        return 'Headset'
    if 'speaker' in d:
        return 'Speakers'
    return 'Accessory'


def build_logitech():
    rows = [r for r in LIVE if r['cat'] == 'LOGITECH']
    norm = lambda m: re.sub(r'\D', '', m)
    lt_names = {norm(r['mpn']): r for r in rows if r['code'].startswith('LT')}
    main_mpns = {norm(r['mpn']) for r in rows if not r['code'].startswith('LT')}
    items = []
    for r in rows:
        if r['code'].startswith('LT') and norm(r['mpn']) in main_mpns:
            continue  # same product listed twice by the supplier: keep one
        typ = logitech_type(r['desc'] or (lt_names.get(norm(r['mpn'])) or {}).get('desc', ''))
        it = build_item(r, 'general', 'Accessories', typ, 'Logitech')
        named = re.match(r'^(the |he )?logitec', r['desc'], re.I)
        if not named:
            lt = lt_names.get(norm(r['mpn']))
            if lt:
                it['name'] = short_name(lt, 'Logitech')
            else:
                it['name'] = f'Logitech {typ} (part no. {r["mpn"]})'
                it['check'] = '; '.join(filter(None, [it['check'], 'Supplier gives no model name — confirm the model name for this part number']))
            it['description'] = fix_case(r['desc']) if r['desc'] else ''
            it['short'] = it['name'] + '.'
        items.append(it)
    return items


SHEETS['Logitech'] = ('Logitech — all products', build_logitech())
def label_type(d):
    d = d.lower()
    if re.match(r'^\S+\s+(9v\s+)?(adapt|case|carry)', d):
        return 'Labelling accessory'
    if re.search(r'label(l)?er\b|label printer|labelling machine|tag printer|line printer|handheld|desktop label|p-touch [a-z]-?\d|pt-?[a-z]?\d{3}', d):
        return 'Labelling machine'
    if 'tape' in d:
        return 'Label tape'
    if 'label' in d:
        return 'Labels'
    return 'Labelling accessory'


SHEETS['Duracell'] = ('Duracell — all products', [build_item(r, 'general', 'Batteries', 'Charger' if 'charger' in r['desc'].lower() else 'Battery', 'Duracell') for r in LIVE if r['cat'] == 'DURACELL'])
SHEETS['Labelling'] = ('Labelling machines — all products & accessories',
                       [build_item(r, 'general', 'Labelling Machines', label_type(r['desc']))
                        for r in LIVE if r['cat'] == 'LABELLING MACHINES & ACCESSORIES'])


def sandisk_cat(d):
    d = d.lower()
    if 'ssd' in d:
        return 'SSD Drives', 'SSD'
    if 'micro' in d or ' sd' in d or 'sdxc' in d or 'sdhc' in d or 'card' in d:
        return 'Memory Cards', 'Memory card'
    if 'usb' in d or 'flash' in d or 'drive' in d:
        return 'USB Flash Drives', 'USB flash drive'
    return 'Hard Drives', 'Portable hard drive'


SHEETS['SanDisk'] = ('SanDisk — all products', [build_item(r, 'general', *sandisk_cat(r['desc']), 'SanDisk') for r in LIVE if r['cat'] == 'SANDISK'])

# ---------------------------------------------------------------- New in the market (hand-written content)
NEW_ITEMS = [
    ('C1643F', 'Canon imageFORCE C1643F A4 Mono Laser Multifunction Printer', 'Printers', 'Canon',
     'New-generation Canon imageFORCE A4 mono laser multifunction printer for busy offices.',
     'The Canon imageFORCE C1643F is part of Canon\'s new imageFORCE range of A4 office devices. It is a monochrome laser multifunction printer for offices that need reliable black-and-white printing, copying and scanning from one compact machine.',
     [('Type', 'A4 monochrome laser multifunction printer'), ('Range', 'Canon imageFORCE'), ('Manufacturer code', '7064C004AA')]),
    ('IFC1333', 'Canon imageFORCE C1333 A4 Colour Laser Multifunction Printer', 'Printers', 'Canon',
     'New Canon imageFORCE A4 colour laser multifunction printer for small and medium offices.',
     'The Canon imageFORCE C1333 brings colour laser printing, copying and scanning to small and medium offices in one compact A4 device from Canon\'s new imageFORCE range.',
     [('Type', 'A4 colour laser multifunction printer'), ('Range', 'Canon imageFORCE'), ('Manufacturer code', '7185C002AA')]),
    ('CPG495', 'Canon PG-495 Black Ink Cartridge', INK, 'Canon',
     'New genuine Canon PG-495 black ink cartridge.',
     'Genuine Canon PG-495 black ink cartridge for sharp, everyday text printing on compatible Canon printers. Ask us to confirm your printer model.',
     [('Type', 'Ink cartridge'), ('Colour', 'Black'), ('Manufacturer code', '7169C001AA')]),
    ('CPG495XL', 'Canon PG-495XL High Yield Black Ink Cartridge', INK, 'Canon',
     'New genuine Canon PG-495XL high-yield black ink cartridge.',
     'The high-yield XL version of the Canon PG-495 black cartridge prints more pages per cartridge, for lower cost per page on compatible Canon printers.',
     [('Type', 'Ink cartridge — high yield (XL)'), ('Colour', 'Black'), ('Manufacturer code', '7168C001AA')]),
    ('CCEXV65BK', 'Canon C-EXV 65 Black Toner', INK, 'Canon',
     'New genuine Canon C-EXV 65 black toner for Canon office copiers.',
     'Genuine Canon C-EXV 65 black toner for compatible Canon office copiers. Ask us to confirm your copier model.',
     [('Type', 'Copier toner'), ('Colour', 'Black'), ('Manufacturer code', '5761C001AA')]),
    ('8011876', 'Kodak E1030 Document Scanner', 'Scanners', 'Kodak',
     'Compact business document scanner — 30 pages per minute, 80-sheet feeder.',
     'The Kodak E1030 is a compact, reliable business-class document scanner for daily office workflows. It scans up to 30 pages / 60 images per minute, holds 80 sheets in the automatic document feeder and detects double-feeds ultrasonically. Includes TWAIN, ISIS and WIA drivers with Smart Touch software.',
     [('Speed', '30 ppm / 60 ipm'), ('Feeder', '80-sheet ADF'), ('Optical resolution', '600 dpi'), ('Daily volume', 'Up to 4 000 pages'),
      ('Connection', 'USB 3.2 Gen 1x1 (USB 2.0 compatible)'), ('Operating systems', 'Windows 10/11, Windows Server, Linux Ubuntu'),
      ('Paper handling', 'Ultrasonic multi-feed detection, straight paper path option'), ('Warranty', '3-year carry-in')]),
    ('8011892', 'Kodak E1040 Document Scanner', 'Scanners', 'Kodak',
     'Fast desktop document scanner — 40 pages per minute, handles ID and embossed cards.',
     'The Kodak E1040 is a fast, reliable desktop sheet-fed scanner that streamlines business document processing. It scans up to 40 pages / 80 images per minute in colour, greyscale or black and white, and handles documents up to 3 m long as well as ID and embossed cards.',
     [('Speed', '40 ppm / 80 ipm at 200/300 dpi'), ('Feeder', '80-sheet ADF'), ('Daily volume', 'Up to 5 000 pages'),
      ('Document size', '51 × 52 mm up to 3 000 mm long'), ('Connection', 'USB 3.2 Gen 1x1 (USB 2.0 compatible)'),
      ('Operating systems', 'Windows 10/11, Windows Server, Linux (Ubuntu, Debian)'), ('Warranty', '3-year carry-in')]),
    ('1014968', 'Kodak S2050 Document Scanner', 'Scanners', 'Kodak',
     'High-speed desktop document scanner — 50 pages per minute for medium-volume offices.',
     'The Kodak S2050 is a compact, high-speed desktop document scanner for efficient office workflows and medium-volume daily scanning, with TWAIN, ISIS and WIA drivers and bundled Smart Touch software.',
     [('Speed', '50 ppm / 100 ipm at 200/300 dpi'), ('Feeder', '80-sheet ADF'), ('Optical resolution', '600 dpi'),
      ('Daily volume', 'Up to 5 000 pages'), ('Connection', 'USB 3.1 Gen 1 (USB 2.0 compatible)'), ('Warranty', '3-year carry-in')]),
    ('8019895', 'Kodak N2050 Network Document Scanner', 'Scanners', 'Kodak',
     'Network-ready A4 document scanner — 50 pages per minute, 100-sheet feeder.',
     'The Kodak N2050 is an A4 network-ready desktop document scanner for branch offices and multi-site organisations. It connects by Gigabit Ethernet or USB and scans up to 50 pages / 100 images per minute.',
     [('Speed', '50 ppm / 100 ipm'), ('Feeder', '100-sheet ADF'), ('Daily volume', 'Up to 8 000 pages'),
      ('Connection', 'Gigabit Ethernet, USB 3.2 Gen 1x1'), ('Warranty', '3-year carry-in')]),
    ('8X220AA', 'HP Poly Blackwire 3320 USB-A Stereo Headset', 'Accessories', 'HP',
     'Wired stereo headset, Microsoft Teams certified, with USB-C to USB-A adapter.',
     'The HP Poly Blackwire 3320 is a comfortable wired stereo headset for calls and online meetings. It is Microsoft Teams certified and comes with a USB-C to USB-A adapter.',
     [('Type', 'Wired stereo headset'), ('Certification', 'Microsoft Teams'), ('Included', 'USB-C to USB-A adapter'), ('Warranty', '2 years'), ('Manufacturer code', '8X220AA')]),
    ('BE1Q0UT', 'HP 495C Multi-Device Dual-Mode Keyboard and Mouse Combo', 'Accessories', 'HP',
     'Multi-device keyboard and mouse combo that switches between your devices.',
     'The HP 495C combo pairs a keyboard and mouse that work with more than one device, so you can switch between your laptop, desktop or tablet.',
     [('Type', 'Keyboard and mouse combo'), ('Connection', 'Dual mode, multi-device'), ('Manufacturer code', 'BE1Q0UT')]),
    ('9X472UT', 'HP Thunderbolt 4 100W G6 Dock', 'Accessories', 'HP',
     'Thunderbolt 4 docking station with up to 100 W laptop charging.',
     'The HP Thunderbolt 4 G6 dock connects your laptop to monitors, network and accessories with a single cable, and charges it at up to 100 W.',
     [('Type', 'Docking station'), ('Connection', 'Thunderbolt 4'), ('Power delivery', 'Up to 100 W'), ('Colour', 'Silver'), ('Manufacturer code', '9X472UT')]),
    ('1B065AA', 'HP E14 G4 14" Portable Monitor', 'Monitors', 'HP',
     '14-inch Full HD portable monitor with two USB-C ports and 5-year warranty.',
     'The HP E14 G4 is a slim 14-inch Full HD portable monitor that adds a second screen to your laptop anywhere, powered and connected through USB-C.',
     [('Screen', '14" Full HD'), ('Refresh rate', '50 Hz'), ('Ports', '2 × USB-C'), ('Colour', 'Silver'), ('Warranty', '5-year next business day'), ('Manufacturer code', '1B065AA')]),
    ('2Z8P3AA', 'HP Prelude 15.6" Laptop Backpack', 'Laptop Bags', 'HP',
     'HP Prelude backpack for laptops up to 15.6 inches.',
     'The HP Prelude backpack carries and protects laptops up to 15.6 inches, with room for chargers and everyday essentials.',
     [('Type', 'Laptop backpack'), ('Fits laptops', 'Up to 15.6"'), ('Manufacturer code', '2Z8P3AA')]),
    ('C78E7AT', 'HP 250R G10 15.6" Notebook — Core 5 120U, 8GB, 512GB SSD, Windows 11 Pro', 'Laptops', 'HP',
     'Business notebook with Intel Core 5 120U, 8 GB RAM, 512 GB SSD and Windows 11 Pro.',
     'The HP 250R G10 is a dependable 15.6-inch Full HD business notebook with an Intel Core 5 120U processor, 8 GB DDR4 memory, a fast 512 GB SSD and Windows 11 Pro.',
     [('Screen', '15.6" Full HD'), ('Processor', 'Intel Core 5 120U'), ('Memory', '8 GB DDR4'), ('Storage', '512 GB SSD'),
      ('Operating system', 'Windows 11 Pro'), ('Colour', 'Silver'), ('Warranty', '1-year carry-in'), ('Manufacturer code', 'C78E7AT')]),
    ('IMC400BLK', 'Ricoh IM C400/C401 Black Toner', INK, 'Ricoh',
     'Genuine Ricoh black toner for IM C400 and IM C401 multifunction printers.',
     'Genuine Ricoh black toner for the Ricoh IM C400 and IM C401 colour multifunction printers.',
     [('Type', 'Toner cartridge'), ('Colour', 'Black'), ('Compatible printers', 'Ricoh IM C400, IM C401'), ('Manufacturer code', '842605')]),
    ('NOVD520BT', 'Novaro DT520BT Thermal Shipping Label Printer', 'Labelling Machines', 'Novaro',
     'Bluetooth thermal printer for shipping labels and waybills.',
     'The Novaro DT520BT is a direct thermal printer for shipping labels and waybills — no ink or toner needed. Ideal for online sellers and dispatch desks.',
     [('Type', 'Thermal shipping label printer'), ('Connection', 'Bluetooth'), ('Labels', 'Shipping labels, e.g. 100 × 150 mm')]),
    ('960001055', 'Logitech C920 HD Pro Webcam', 'Accessories', 'Logitech',
     'Full HD 1080p webcam with dual stereo microphones and automatic light correction.',
     'The Logitech C920 HD Pro is a premium Full HD webcam for video calls, streaming and content creation, with dual stereo microphones, automatic light correction and a universal clip for laptops, monitors and tripods.',
     [('Video', 'Full HD 1080p'), ('Microphones', 'Dual stereo'), ('Features', 'Automatic light correction, universal mounting clip'), ('Connection', 'USB'), ('Manufacturer code', '960-001055')]),
    ('B1503CVAC73210G0X', 'Asus ExpertBook B1 15.6" Laptop — 32GB RAM, 1TB SSD, Windows 11 Pro', 'Laptops', 'Asus',
     'Business laptop with 32 GB memory, 1 TB SSD and Windows 11 Pro.',
     'The Asus ExpertBook B1 is a 15.6-inch business laptop with 32 GB memory, a 1 TB SSD and Windows 11 Pro — plenty of power for multitasking at work.',
     [('Screen', '15.6"'), ('Memory', '32 GB'), ('Storage', '1 TB SSD'), ('Operating system', 'Windows 11 Pro'), ('Manufacturer code', '90NX0801-M09AL0')]),
    ('M1X', 'ViewSonic M1X Smart LED Portable Projector with Harman Kardon Speaker', 'Projectors', 'ViewSonic',
     'Portable smart LED projector with built-in Harman Kardon speaker.',
     'The ViewSonic M1X is a compact smart LED projector with a built-in Harman Kardon speaker, for presentations and movies wherever you go.',
     [('Type', 'Portable LED projector'), ('Audio', 'Harman Kardon speaker'), ('Manufacturer code', 'M1X')]),
]
new_items = []
for code, name, cat, brand, short, desc, specs in NEW_ITEMS:
    r = BY_CODE[code]
    it = build_item(r, 'general', cat, '', brand)
    it.update(name=name, short=short, description=desc, specs=specs, new=True, type=dict(specs).get('Type', ''))
    new_items.append(it)

# ---------------------------------------------------------------- write workbook
NAVY, GOLD, PALE = '1A2F45', 'C9A84C', 'F7EFD9'
F = 'Arial'
HFONT = Font(name=F, bold=True, color='FFFFFF', size=10)
HFILL = PatternFill('solid', fgColor=NAVY)
BODY = Font(name=F, size=10)
BLUE = Font(name=F, size=10, color='0000FF')
BOLD = Font(name=F, size=10, bold=True)
THIN = Border(bottom=Side(style='thin', color='E3E7EC'))
wrap = Alignment(wrap_text=True, vertical='top')
top = Alignment(vertical='top')

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Summary'
ws['A1'] = 'Setwel Africa — product catalogue from supplier price list (effective 2 Oct 2026)'
ws['A1'].font = Font(name=F, size=14, bold=True, color=NAVY)
ws['A3'] = 'Markup on supplier cost (excl. VAT)'
ws['A3'].font = BOLD
ws['C3'] = MARKUP
ws['C3'].number_format = '0%'
ws['C3'].font = Font(name=F, size=11, bold=True, color='0000FF')
ws['C3'].fill = PatternFill('solid', fgColor='FFFF00')
ws['D3'] = '← change this one cell and every selling price in this workbook updates (55% = Setwel\'s markup, which already covers the 15% VAT paid to the supplier)'
ws['D3'].font = Font(name=F, size=9, italic=True, color='5B6573')
ws['A4'] = 'Selling price rule'
ws['A4'].font = BOLD
ws['C4'] = 'Supplier cost × (1 + markup), rounded UP to the next rand — same rule as the website'
ws['C4'].font = BODY
MARK = 'Summary!$C$3'

ws['A6'], ws['B6'], ws['C6'], ws['D6'] = 'Sheet', 'Products', 'What is in it', 'Website category'
for c in 'ABCD':
    ws[f'{c}6'].font = HFONT
    ws[f'{c}6'].fill = HFILL
row = 7

COLS = ['SKU (supplier code)', 'Manufacturer code', 'Product name', 'Brand', 'Type', 'Colour', 'Page yield', 'Capacity',
        'Compatible printers', 'Supplier cost excl. VAT (R)', 'Setwel selling price (R)', 'Barcode', 'Short description',
        'Full description', 'Specifications', 'Photo file name (for Bulk images)', 'Please check', 'Supplier description (original)']
WIDTHS = [16, 16, 46, 11, 18, 10, 11, 10, 38, 13, 13, 16, 44, 55, 38, 20, 30, 55]
WRAP_COLS = {3, 9, 13, 14, 15, 17, 18}


def write_sheet(name, title, items):
    sh = wb.create_sheet(name)
    sh['A1'] = title
    sh['A1'].font = Font(name=F, size=13, bold=True, color=NAVY)
    sh['A2'] = (f'{len(items)} products · discontinued (EOL) items excluded · selling price = supplier cost + markup on the Summary sheet, '
                'rounded up to the next rand · blue = supplier input')
    sh['A2'].font = Font(name=F, size=9, italic=True, color='5B6573')
    for i, h in enumerate(COLS, 1):
        c = sh.cell(row=4, column=i, value=h)
        c.font, c.fill, c.alignment = HFONT, HFILL, Alignment(wrap_text=True, vertical='center')
    for i, it in enumerate(items):
        r = 5 + i
        spec_txt = '\n'.join(f'{k}: {v}' for k, v in it['specs'])
        vals = [it['sku'], it['mpn'], it['name'], it['brand'], it['type'], it['colour'], it['yield'], it['capacity'],
                it['compatible'], it['cost'], f'=IF(J{r}="","Price on request",CEILING(J{r}*(1+{MARK}),1))', it['barcode'],
                it['short'], it.get('description') or '', spec_txt, f"{it['sku']}.jpg", it.get('check', ''), it['source']]
        for j, v in enumerate(vals, 1):
            c = sh.cell(row=r, column=j, value=v if v != '' else None)
            c.font = BLUE if j == 10 else BODY
            c.alignment = wrap if j in WRAP_COLS else top
            c.border = THIN
        sh.cell(row=r, column=10).number_format = '#,##0.00'
        sh.cell(row=r, column=11).number_format = '#,##0'
        sh.cell(row=r, column=11).font = BOLD
        if it.get('check'):
            sh.cell(row=r, column=17).fill = PatternFill('solid', fgColor='FFF4CC')
    for i, w in enumerate(WIDTHS, 1):
        sh.column_dimensions[get_column_letter(i)].width = w
    sh.freeze_panes = 'D5'
    sh.auto_filter.ref = f'A4:{get_column_letter(len(COLS))}{4 + len(items)}'
    return sh


all_for_import = []
for key, (title, items) in SHEETS.items():
    write_sheet(key, title, items)
    ws.cell(row=row, column=1, value=key).font = BOLD
    ws.cell(row=row, column=2, value=f"=COUNTA('{key}'!A5:A{4 + len(items)})").font = BODY
    ws.cell(row=row, column=3, value=title).font = BODY
    ws.cell(row=row, column=4, value=', '.join(sorted({i['category'] for i in items}))).font = BODY
    row += 1
    all_for_import += items
write_sheet('New in Market', 'New in the market — 20 hand-picked new products with full descriptions and specs', new_items)
ws.cell(row=row, column=1, value='New in Market').font = BOLD
ws.cell(row=row, column=2, value=f"=COUNTA('New in Market'!A5:A{4 + len(new_items)})").font = BODY
ws.cell(row=row, column=3, value='20 new products with full descriptions and specs').font = BODY
ws.cell(row=row, column=4, value=', '.join(sorted({i['category'] for i in new_items}))).font = BODY
row += 1
all_for_import += new_items

# de-duplicate (New in Market wins)
seen = OrderedDict()
for it in all_for_import:
    seen[it['sku']] = it
imp_items = list(seen.values())
ws.cell(row=row, column=1, value='Website import').font = BOLD
ws.cell(row=row, column=2, value=f"=COUNTA('Website import'!A2:A{1 + len(imp_items)})").font = BODY
ws.cell(row=row, column=3, value='Everything above in one sheet, in the store\'s import format (duplicates removed)').font = BODY
row += 2

notes = [
    'HOW TO PUT THESE PRODUCTS ON THE WEBSITE',
    '1. Admin → Import / update prices → upload the separate file "Setwel-Website-Import.xlsx" (it holds only the Website import sheet).',
    '   Edit products in that file (or in the Website import sheet here) before uploading if you want to change names or descriptions.',
    '2. Matching is automatic. Leave "Selling price" empty in the import — the website adds your markup to the supplier cost itself.',
    '3. Tick "Add NEW products as hidden" if you want to add photos before customers see them.',
    '4. Photos: save each picture with the file name shown in the "Photo file name" column (SKU.jpg) and upload them all at Admin → Bulk images.',
    '',
    'HOW THE PRODUCTS WERE CHOSEN',
    '• "Newest" for HP, Canon and Kyocera = the newest cartridge series for the newest printer models (e.g. HP 145/151/153/154/222/230/225 for',
    '  2023 LaserJet Pro, Tank and Colour printers; HP 925/938 for 2024 OfficeJet Pro; Canon CRG-075/071/070/069; Canon GI-46/45/43; Kyocera',
    '  TK-3400/7135/8375). The supplier list has no launch dates, so this is based on printer model generations — please check it suits your customers.',
    '• "New in Market" = items the supplier marks NEW on this price list, plus well-known current products. Warranty packs and very large',
    '  copiers/enterprise printers were left out.',
    '• Discontinued (EOL) items are excluded everywhere.',
    '',
    'PLEASE CHECK BEFORE GOING LIVE',
    '• Product names, specs and compatible printers were cleaned up automatically from the supplier descriptions. Check a few in each sheet.',
    '• Specs for the Canon imageFORCE printers and some accessories are limited to what the supplier list says — add more if you have the brochure.',
    '• ViewSonic M1X: the supplier list calls it a "portable monitor", but ViewSonic M1-series products are portable projectors. It is listed',
    '  here as a projector — please confirm with the supplier.',
    '• Photos are not included: the supplier price list has no pictures. Use the distributor\'s white-label flyers or the manufacturer\'s reseller images.',
]
for n in notes:
    c = ws.cell(row=row, column=1, value=n)
    c.font = Font(name=F, size=10, bold=n.isupper() and n != '')
    row += 1
ws.column_dimensions['A'].width = 20
ws.column_dimensions['B'].width = 10
ws.column_dimensions['C'].width = 62
ws.column_dimensions['D'].width = 40

# ---- Website import sheet (exact importer headings, values only)
imp = wb.create_sheet('Website import')
heads = ['SKU', 'Name', 'Brand', 'Category', 'Supplier cost (excl VAT)', 'Selling price', 'Sale price', 'Sale ends', 'Stock status',
         'Visible', 'Featured', 'Short description', 'Description', 'Specifications', 'Compatible printers', 'Image URLs', 'MPN', 'GTIN',
         'Warranty', 'SEO title', 'SEO description']
for i, h in enumerate(heads, 1):
    c = imp.cell(row=1, column=i, value=h)
    c.font, c.fill = HFONT, HFILL
for i, it in enumerate(imp_items, 2):
    desc = it.get('description') or it['short']
    spec_txt = ' | '.join(f'{k}: {v}' for k, v in it['specs'])
    gtin = re.sub(r'\D', '', it['barcode'])
    vals = [it['sku'], it['name'], it['brand'], it['category'], it['cost'], None, None, None, 'in stock', 'yes',
            'yes' if it.get('new') else 'no', it['short'], desc, spec_txt, it['compatible'], None, it['mpn'],
            gtin if len(gtin) in (8, 12, 13, 14) else None, next((v for k, v in it['specs'] if k == 'Warranty'), None), None, None]
    for j, v in enumerate(vals, 1):
        c = imp.cell(row=i, column=j, value=v)
        c.font = BODY
    imp.cell(row=i, column=5).number_format = '0.00'
for i, w in enumerate([16, 50, 12, 18, 12, 10, 8, 10, 10, 7, 8, 45, 55, 45, 40, 10, 16, 15, 18, 10, 10], 1):
    imp.column_dimensions[get_column_letter(i)].width = w
imp.freeze_panes = 'C2'
imp.auto_filter.ref = f'A1:U{len(imp_items) + 1}'

wb.save(OUT)

# The same import sheet on its own, so it can be uploaded straight to Admin → Import
single = openpyxl.Workbook()
sw = single.active
sw.title = 'Website import'
for row in imp.iter_rows(values_only=True):
    sw.append(list(row))
for i, h in enumerate(heads, 1):
    sw.cell(row=1, column=i).font, sw.cell(row=1, column=i).fill = HFONT, HFILL
for col, dim in imp.column_dimensions.items():
    sw.column_dimensions[col].width = dim.width
sw.freeze_panes = 'C2'
single_path = OUT.replace('.xlsx', '') .rsplit('/', 1)[0] + '/Setwel-Website-Import.xlsx'
single.save(single_path)
print('Saved', single_path)
print('Saved', OUT)
for k, (t, items) in SHEETS.items():
    print(f'{k:12} {len(items):4}')
print(f'{"New":12} {len(new_items):4}')
print(f'{"Import":12} {len(imp_items):4}')
